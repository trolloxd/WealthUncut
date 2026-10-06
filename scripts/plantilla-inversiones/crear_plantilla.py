"""
Genera la plantilla de seguimiento de inversiones en fondos indexados de WealthUncut (Excel).

Uso (necesita `pip install xlsxwriter openpyxl`):
    python scripts/plantilla-inversiones/crear_plantilla.py <ruta-salida.xlsx> [xlsx-calculado-por-Excel]

El segundo argumento es opcional: un xlsx ya calculado por Excel, del que se copian los resultados de las formulas como
valores en cache (asi las vistas previas sin Excel no muestran las celdas vacias). Flujo completo en CLAUDE.md.
Los fondos y los numeros son de ejemplo (nombres ficticios, no recomendaciones).
"""
import random
import sys
from datetime import date, datetime

import xlsxwriter

RUTA = sys.argv[1] if len(sys.argv) > 1 else 'Plantilla-seguimiento-inversiones-WealthUncut.xlsx'
CACHE = sys.argv[2] if len(sys.argv) > 2 else None

_valores = {}
if CACHE:
    import openpyxl
    from openpyxl.utils.datetime import to_excel
    from xlsxwriter.worksheet import Worksheet as _W

    for _ws in openpyxl.load_workbook(CACHE, data_only=True):
        for _fila in _ws.iter_rows():
            for _c in _fila:
                if _c.value is not None:
                    v = _c.value
                    if hasattr(v, 'year'):
                        v = to_excel(v)
                    _valores[(_ws.title, _c.row - 1, _c.column - 1)] = v
    _orig = _W._write_formula

    def _con_cache(self, row, col, formula, cell_format=None, value=0):
        return _orig(self, row, col, formula, cell_format, _valores.get((self.name, row, col), ''))

    _W._write_formula = _con_cache

VERDE = '#1E4A3A'
VERDE_CLARO = '#E3EDE7'
CREMA = '#FAF6ED'
ROJO = '#B3402A'
ROJO_CLARO = '#F6E3DE'
TINTA = '#1F2A24'
GRIS = '#6B7A72'
BORDE = '#D9D2C0'
AZUL = '#0000FF'
AMARILLO = '#FFF8DC'
DORADO = '#D8A65C'

wb = xlsxwriter.Workbook(RUTA)
wb.set_properties({
    'title': 'Plantilla de seguimiento de inversiones en fondos indexados',
    'subject': 'Cartera, aportaciones, rentabilidad, reequilibrado e impuesto estimado',
    'author': 'David Pérez Mitjà · WealthUncut',
    'company': 'WealthUncut',
    'comments': 'Plantilla gratuita de wealthuncut.com. Información educativa, no asesoramiento financiero.',
    'website': 'https://wealthuncut.com/plantilla-seguimiento-inversiones/',
})
_cache = {}


def F(**kw):
    base = {'font_name': 'Arial', 'font_size': 10, 'font_color': TINTA, 'valign': 'vcenter'}
    base.update(kw)
    k = tuple(sorted(base.items()))
    if k not in _cache:
        _cache[k] = wb.add_format(base)
    return _cache[k]


EUR = '#,##0.00 "€"'
EUR0 = '#,##0 "€"'
EUR_G = '#,##0.00 "€";-#,##0.00 "€";"-"'
PCT = '0%'
PCT1 = '0.0%'
PCT2 = '0.00%'
FECHA = 'dd/mm/yyyy'
PART = '#,##0.0000'

f_titulo = F(font_size=20, bold=True, font_color=VERDE)
f_sub = F(font_color=GRIS, italic=True)
f_h = F(bold=True, font_color='#FFFFFF', bg_color=VERDE, align='center', border=1, border_color=VERDE, text_wrap=True)
f_hizq = F(bold=True, font_color='#FFFFFF', bg_color=VERDE, align='left', border=1, border_color=VERDE, indent=1)
f_celda = F(border=1, border_color=BORDE)
f_celda_c = F(border=1, border_color=BORDE, align='center')
f_eur = F(border=1, border_color=BORDE, num_format=EUR)
f_eur_g = F(border=1, border_color=BORDE, num_format=EUR_G)
f_pct = F(border=1, border_color=BORDE, num_format=PCT1, align='center')
f_part = F(border=1, border_color=BORDE, num_format=PART)
f_in = F(border=1, border_color=BORDE, font_color=AZUL, bg_color=AMARILLO)
f_in_c = F(border=1, border_color=BORDE, font_color=AZUL, bg_color=AMARILLO, align='center')
f_in_eur = F(border=1, border_color=BORDE, font_color=AZUL, bg_color=AMARILLO, num_format=EUR)
f_in_vl = F(border=1, border_color=BORDE, font_color=AZUL, bg_color=AMARILLO, num_format='#,##0.0000 "€"')
f_in_part = F(border=1, border_color=BORDE, font_color=AZUL, bg_color=AMARILLO, num_format=PART)
f_in_pct = F(border=1, border_color=BORDE, font_color=AZUL, bg_color=AMARILLO, num_format=PCT, align='center')
f_in_pct2 = F(border=1, border_color=BORDE, font_color=AZUL, bg_color=AMARILLO, num_format='0.00', align='center')
f_in_fecha = F(border=1, border_color=BORDE, font_color=AZUL, bg_color=AMARILLO, num_format=FECHA, align='center')
f_tot = F(bold=True, border=1, border_color=BORDE, bg_color=VERDE_CLARO)
f_tot_eur = F(bold=True, border=1, border_color=BORDE, bg_color=VERDE_CLARO, num_format=EUR)
f_tot_pct = F(bold=True, border=1, border_color=BORDE, bg_color=VERDE_CLARO, num_format=PCT1, align='center')
f_aux = F(font_color=GRIS, font_size=9)
f_aux_eur = F(font_color=GRIS, font_size=9, num_format=EUR)
f_aux_fecha = F(font_color=GRIS, font_size=9, num_format=FECHA)
f_nota = F(font_color=GRIS, font_size=9, italic=True, text_wrap=True, valign='top')
f_nota1 = F(font_color=GRIS, font_size=9, italic=True)
f_link = F(font_color='#0563C1', underline=1)


def hoja(nombre, color):
    ws = wb.add_worksheet(nombre)
    ws.set_tab_color(color)
    ws.hide_gridlines(2)
    ws.set_column('A:A', 2, F())
    ws.set_default_row(18)
    ws.set_row(0, 32)
    return ws


ws_i = hoja('Instrucciones', VERDE)
ws_r = hoja('Resumen', VERDE)
ws_f = hoja('Fondos', ROJO)
ws_o = hoja('Operaciones', ROJO)
ws_v = hoja('Valoraciones', GRIS)
ws_p = hoja('Proyección', GRIS)
ws_l = hoja('Listas', BORDE)
ws_c = hoja('Calculo', BORDE)
ws_c.hide()

TIPOS_ACTIVO = ['Renta variable', 'Renta fija', 'Monetario', 'Otros']
TIPOS_OP = ['Compra', 'Venta', 'Traspaso entrada', 'Traspaso salida']
R0, R1 = 6, 505  # filas de Operaciones
F0, F1 = 5, 12   # filas de Fondos (8 fondos)
V0, V1 = 5, 64   # filas de Valoraciones

OP = {c: f'Operaciones!${c}${R0}:${c}${R1}' for c in 'BCDEFGHIJKL'}

# ------------------------------------------------------------- Listas (tipos y tramos del IRPF del ahorro)
ws_l.set_column('B:B', 18)
ws_l.set_column('D:D', 18)
ws_l.set_column('F:H', 14)
ws_l.write('B1', 'Tipos de activo', f_hizq)
ws_l.write('D1', 'Tipos de operación', f_hizq)
for i, t in enumerate(TIPOS_ACTIVO):
    ws_l.write(1 + i, 1, t, f_celda)
for i, t in enumerate(TIPOS_OP):
    ws_l.write(1 + i, 3, t, f_celda)
ws_l.write('F1', 'Desde (€)', f_h)
ws_l.write('G1', 'Hasta (€)', f_h)
ws_l.write('H1', 'Tipo', f_h)
ws_l.write('I1', 'Base del tramo', f_h)
TRAMOS = [(0, 6000, 0.19), (6000, 50000, 0.21), (50000, 200000, 0.23), (200000, 300000, 0.27), (300000, 1e12, 0.30)]
for i, (d, h, t) in enumerate(TRAMOS):
    r = 2 + i
    ws_l.write_number(f'F{r}', d, f_in_eur)
    ws_l.write_number(f'G{r}', h, f_in_eur)
    ws_l.write_number(f'H{r}', t, f_in_pct)
    ws_l.write_formula(f'I{r}', f'=MAX(0,MIN(Resumen!$E$15,G{r})-F{r})', f_eur)
ws_l.write('F8', 'Tramos del IRPF del ahorro (estatal y autonómico, régimen común). Fuente: AEAT, Manual práctico de Renta 2025. Verificados el 23/09/2026.', f_nota1)
ws_l.write('F9', 'Si cambian, edítalos aquí. Más información en wealthuncut.com/blog/tramos-irpf-ahorro/', f_nota1)
ws_l.write('B10', 'Esta hoja alimenta los desplegables y el cálculo del impuesto.', f_nota1)

# ------------------------------------------------------------- Operaciones
ws_o.set_column('B:B', 12)
ws_o.set_column('C:C', 30)
ws_o.set_column('D:D', 17)
ws_o.set_column('E:F', 16)
ws_o.set_column('G:G', 12)
ws_o.set_column('H:I', 16)
ws_o.set_column('J:J', 26)
ws_o.set_column('K:L', 14, f_aux)
ws_o.write('B1', 'Operaciones', f_titulo)
ws_o.write('B2', 'Apunta cada compra, venta o traspaso de fondos. Con las participaciones y el valor liquidativo de la orden, el importe sale solo.', f_sub)
ws_o.write('B3', 'Operaciones registradas:', F(bold=True))
ws_o.write_formula('D3', f'=COUNT(B{R0}:B{R1})', F(bold=True, num_format='0', align='left'))
cab = ['Fecha', 'Fondo', 'Tipo', 'Participaciones', 'Valor liquidativo (€)', 'Comisión (€)', 'Importe (auto)', 'Coste heredado (€)', 'Nota', 'Flujo (auto)', 'Fecha flujo']
for i, h in enumerate(cab):
    ws_o.write(4, 1 + i, h, f_h if i < 9 else F(font_color=GRIS, font_size=9, align='center'))
ws_o.set_row(4, 30)

# --- datos de ejemplo (fondos y cifras ficticios) ---
FONDOS = [
    ('Indexado mundo desarrollado (ejemplo)', 'EJEMPLO-001', 'Renta variable', 0.12, 0.60, 10.00, 0.0085, 0.025),
    ('Indexado mercados emergentes (ejemplo)', 'EJEMPLO-002', 'Renta variable', 0.23, 0.15, 8.00, 0.0060, 0.030),
    ('Indexado bonos euro (ejemplo)', 'EJEMPLO-003', 'Renta fija', 0.15, 0.20, 12.00, 0.0015, 0.008),
    ('Fondo monetario (ejemplo)', 'EJEMPLO-004', 'Monetario', 0.10, 0.05, 100.00, 0.0025, 0.0006),
]
random.seed(7)
vl = {}  # (indice_fondo, año, mes) -> VL
estado = [f[5] for f in FONDOS]
meses = [(2025, m) for m in range(1, 13)] + [(2026, m) for m in range(1, 10)]
for (a, m) in meses:
    for k, fnd in enumerate(FONDOS):
        estado[k] *= 1 + fnd[6] + random.uniform(-fnd[7], fnd[7])
        vl[(k, a, m)] = round(estado[k], 4)
OPS = []
MENSUAL = [120, 40, 40]
for (a, m) in meses:
    for k in range(3):
        OPS.append([date(a, m, 5), k, 'Compra', round(MENSUAL[k] / vl[(k, a, m)], 4), vl[(k, a, m)], 0.0, None, ''])
OPS.insert(0, [date(2025, 1, 3), 3, 'Compra', round(600 / 100.0, 4), 100.0, 0.0, None, 'Colchón inicial'])
# una venta parcial del monetario y un traspaso de renta fija a renta variable
OPS.append([date(2026, 3, 20), 3, 'Venta', 2.0, vl[(3, 2026, 3)], 0.0, None, 'Retiro para un gasto'])
sal_part = 8.0
OPS.append([date(2026, 6, 15), 2, 'Traspaso salida', sal_part, vl[(2, 2026, 6)], 0.0, None, 'Traspaso a otro fondo'])
OPS.sort(key=lambda x: (x[0], x[1]))


def coste_medio_hasta(k, hasta):
    comp = [o for o in OPS if o[1] == k and o[2] == 'Compra' and o[0] <= hasta]
    pa = sum(o[3] for o in comp)
    ca = sum(o[3] * o[4] + o[5] for o in comp)
    return ca / pa if pa else 0


coste_heredado = round(coste_medio_hasta(2, date(2026, 6, 15)) * sal_part, 2)
entrada_part = round(sal_part * vl[(2, 2026, 6)] / vl[(1, 2026, 6)], 4)
OPS.append([date(2026, 6, 15), 1, 'Traspaso entrada', entrada_part, vl[(1, 2026, 6)], 0.0, coste_heredado, 'Viene del fondo de bonos'])
OPS.sort(key=lambda x: (x[0], x[1], 0 if x[2] != 'Traspaso entrada' else 1))

for r in range(R0, R1 + 1):
    fila = r - 1
    i = r - R0
    if i < len(OPS):
        d, k, tipo, part, precio, com, her, nota = OPS[i]
        ws_o.write_datetime(fila, 1, datetime(d.year, d.month, d.day), f_in_fecha)
        ws_o.write(fila, 2, FONDOS[k][0], f_in)
        ws_o.write(fila, 3, tipo, f_in_c)
        ws_o.write_number(fila, 4, part, f_in_part)
        ws_o.write_number(fila, 5, precio, f_in_vl)
        ws_o.write_number(fila, 6, com, f_in_eur)
        if her is None:
            ws_o.write_blank(fila, 8, None, f_in_eur)
        else:
            ws_o.write_number(fila, 8, her, f_in_eur)
        ws_o.write(fila, 9, nota, f_in)
    else:
        ws_o.write_blank(fila, 1, None, f_in_fecha)
        ws_o.write_blank(fila, 2, None, f_in)
        ws_o.write_blank(fila, 3, None, f_in_c)
        ws_o.write_blank(fila, 4, None, f_in_part)
        ws_o.write_blank(fila, 5, None, f_in_vl)
        ws_o.write_blank(fila, 6, None, f_in_eur)
        ws_o.write_blank(fila, 8, None, f_in_eur)
        ws_o.write_blank(fila, 9, None, f_in)
    ws_o.write_formula(fila, 7, f'=IF(OR($E{r}="",$F{r}=""),"",$E{r}*$F{r})', f_eur)
    ws_o.write_formula(fila, 10, f'=IF($B{r}="","",IF($D{r}="Compra",-(N(H{r})+N(G{r})),IF($D{r}="Venta",N(H{r})-N(G{r}),0)))', f_aux_eur)
    ws_o.write_formula(fila, 11, f'=IF($B{r}="","",$B{r})', f_aux_fecha)
ws_o.data_validation(f'B{R0}:B{R1}', {'validate': 'date', 'criteria': '>=', 'value': date(2000, 1, 1), 'error_title': 'Fecha no válida', 'error_message': 'Escribe una fecha, por ejemplo 05/09/2026.'})
ws_o.data_validation(f'C{R0}:C{R1}', {'validate': 'list', 'source': f'=Fondos!$B${F0}:$B${F1}', 'error_title': 'Fondo no encontrado', 'error_message': 'Elige un fondo de la lista o añádelo antes en la hoja Fondos.'})
ws_o.data_validation(f'D{R0}:D{R1}', {'validate': 'list', 'source': '=Listas!$D$2:$D$5'})
ws_o.data_validation(f'E{R0}:E{R1}', {'validate': 'decimal', 'criteria': '>=', 'value': 0, 'error_message': 'Escribe un número positivo.'})
ws_o.data_validation(f'F{R0}:F{R1}', {'validate': 'decimal', 'criteria': '>=', 'value': 0, 'error_message': 'Escribe un número positivo.'})
ws_o.data_validation(f'G{R0}:G{R1}', {'validate': 'decimal', 'criteria': '>=', 'value': 0, 'error_message': 'Escribe un número positivo.'})
ws_o.data_validation(f'I{R0}:I{R1}', {'validate': 'decimal', 'criteria': '>=', 'value': 0, 'ignore_blank': True, 'error_message': 'Escribe un número positivo.'})
ws_o.conditional_format(f'D{R0}:D{R1}', {'type': 'text', 'criteria': 'begins with', 'value': 'Venta', 'format': F(font_color=ROJO, bold=True, bg_color=ROJO_CLARO)})
ws_o.conditional_format(f'D{R0}:D{R1}', {'type': 'text', 'criteria': 'begins with', 'value': 'Traspaso', 'format': F(font_color='#8A6D00', bold=True, bg_color='#FFF3C4')})
ws_o.autofilter(f'B5:J{R1}')
ws_o.freeze_panes(5, 0)
ws_o.set_landscape()
ws_o.fit_to_pages(1, 1)
ws_o.print_area(f'B1:J{R0 + 34}')

# ------------------------------------------------------------- Fondos
ws_f.set_column('B:B', 36)
ws_f.set_column('C:C', 14)
ws_f.set_column('D:D', 15)
ws_f.set_column('E:F', 11)
ws_f.set_column('G:G', 15)
ws_f.set_column('H:H', 17)
ws_f.set_column('I:K', 15)
ws_f.set_column('L:N', 12)
ws_f.set_column('O:P', 15)
ws_f.set_column('R:S', 14, f_aux)
ws_f.write('B1', 'Tus fondos', f_titulo)
ws_f.write('B2', 'Una fila por fondo. Rellena el nombre, el tipo, el TER, el peso que quieres y el valor liquidativo de hoy: lo demás se calcula con tus operaciones.', f_sub)
cab = ['Fondo', 'ISIN', 'Tipo de activo', 'TER (%)', 'Peso objetivo', 'Participaciones', 'Valor liquidativo hoy (€)', 'Valor actual (€)',
       'Coste de lo que tienes (€)', 'Ganancia latente (€)', 'Rentabilidad', 'Peso actual', 'Desviación', 'Ajuste para reequilibrar (€)', 'Coste anual por TER (€)']
for i, h in enumerate(cab):
    ws_f.write(3, 1 + i, h, f_h)
ws_f.write(3, 17, 'Part. adquiridas', F(font_color=GRIS, font_size=9, align='center'))
ws_f.write(3, 18, 'Coste adquirido', F(font_color=GRIS, font_size=9, align='center'))
ws_f.set_row(3, 42)
ULT = {k: vl[(k, 2026, 9)] for k in range(4)}
for r in range(F0, F1 + 1):
    fila = r - 1
    i = r - F0
    if i < len(FONDOS):
        n, isin, tipo, ter, obj, _, _, _ = FONDOS[i]
        ws_f.write(fila, 1, n, f_in)
        ws_f.write(fila, 2, isin, f_in_c)
        ws_f.write(fila, 3, tipo, f_in_c)
        ws_f.write_number(fila, 4, ter, f_in_pct2)
        ws_f.write_number(fila, 5, obj, f_in_pct)
        ws_f.write_number(fila, 7, ULT[i], f_in_vl)
    else:
        for c, fm in ((1, f_in), (2, f_in_c), (3, f_in_c), (4, f_in_pct2), (5, f_in_pct), (7, f_in_vl)):
            ws_f.write_blank(fila, c, None, fm)
    ws_f.write_formula(fila, 6, f'=IF($B{r}="","",SUMIFS({OP["E"]},{OP["C"]},$B{r},{OP["D"]},"Compra")+SUMIFS({OP["E"]},{OP["C"]},$B{r},{OP["D"]},"Traspaso entrada")-SUMIFS({OP["E"]},{OP["C"]},$B{r},{OP["D"]},"Venta")-SUMIFS({OP["E"]},{OP["C"]},$B{r},{OP["D"]},"Traspaso salida"))', f_part)
    ws_f.write_formula(fila, 8, f'=IF($B{r}="","",G{r}*N(H{r}))', f_eur)
    ws_f.write_formula(fila, 9, f'=IF($B{r}="","",IF(R{r}=0,0,S{r}/R{r}*G{r}))', f_eur)
    ws_f.write_formula(fila, 10, f'=IF($B{r}="","",I{r}-J{r})', f_eur_g)
    ws_f.write_formula(fila, 11, f'=IF(OR($B{r}="",N(J{r})=0),"",K{r}/J{r})', f_pct)
    ws_f.write_formula(fila, 12, f'=IF(OR($B{r}="",$I$13=0),"",I{r}/$I$13)', f_pct)
    ws_f.write_formula(fila, 13, f'=IF(OR($B{r}="",M{r}=""),"",M{r}-F{r})', F(border=1, border_color=BORDE, num_format='+0.0%;-0.0%;0.0%', align='center'))
    ws_f.write_formula(fila, 14, f'=IF($B{r}="","",F{r}*$I$13-I{r})', f_eur_g)
    ws_f.write_formula(fila, 15, f'=IF($B{r}="","",I{r}*N(E{r})/100)', f_eur)
    ws_f.write_formula(fila, 17, f'=IF($B{r}="","",SUMIFS({OP["E"]},{OP["C"]},$B{r},{OP["D"]},"Compra")+SUMIFS({OP["E"]},{OP["C"]},$B{r},{OP["D"]},"Traspaso entrada"))', F(font_color=GRIS, font_size=9, num_format=PART))
    ws_f.write_formula(fila, 18, f'=IF($B{r}="","",SUMIFS({OP["H"]},{OP["C"]},$B{r},{OP["D"]},"Compra")+SUMIFS({OP["G"]},{OP["C"]},$B{r},{OP["D"]},"Compra")+SUMIFS({OP["I"]},{OP["C"]},$B{r},{OP["D"]},"Traspaso entrada"))', f_aux_eur)
ws_f.write('B13', 'Total', f_tot)
for c in 'CDEGHL':
    ws_f.write_blank(f'{c}13', None, f_tot)
ws_f.write_formula('F13', f'=SUM(F{F0}:F{F1})', f_tot_pct)
for col in 'IJKOP':
    ws_f.write_formula(f'{col}13', f'=SUM({col}{F0}:{col}{F1})', f_tot_eur)
ws_f.write_formula('L13', '=IF(J13=0,"",K13/J13)', f_tot_pct)
ws_f.write_formula('M13', f'=SUM(M{F0}:M{F1})', f_tot_pct)
ws_f.write_blank('N13', None, f_tot)
ws_f.write_formula('B14', '=IF(ABS(F13-1)>0.0001,"Atención: tus pesos objetivo suman "&TEXT(F13,"0%")&" y deberían sumar 100%.","")', F(font_color=ROJO, bold=True))
ws_f.data_validation(f'D{F0}:D{F1}', {'validate': 'list', 'source': '=Listas!$B$2:$B$5'})
ws_f.data_validation(f'E{F0}:E{F1}', {'validate': 'decimal', 'criteria': 'between', 'minimum': 0, 'maximum': 10, 'error_message': 'Escribe el TER en porcentaje, por ejemplo 0,15.'})
ws_f.data_validation(f'F{F0}:F{F1}', {'validate': 'decimal', 'criteria': 'between', 'minimum': 0, 'maximum': 1, 'error_message': 'Escribe el peso entre 0% y 100%.'})
ws_f.conditional_format(f'K{F0}:K{F1}', {'type': 'cell', 'criteria': '<', 'value': 0, 'format': F(font_color=ROJO, bold=True)})
ws_f.conditional_format(f'L{F0}:L{F1}', {'type': 'cell', 'criteria': '<', 'value': 0, 'format': F(font_color=ROJO, bold=True)})
ws_f.conditional_format(f'N{F0}:N{F1}', {'type': 'formula', 'criteria': f'=AND(ISNUMBER($N{F0}),ABS($N{F0})>0.05)', 'format': F(font_color=ROJO, bg_color=ROJO_CLARO, bold=True)})
ws_f.conditional_format(f'M{F0}:M{F1}', {'type': 'data_bar', 'bar_color': '#7FA392', 'bar_solid': True, 'min_type': 'num', 'min_value': 0, 'max_type': 'num', 'max_value': 1})
ws_f.write('B16', 'Ajuste para reequilibrar: positivo = comprar ese importe; negativo = vender. Sale de: peso objetivo × valor total de la cartera − valor actual.', f_nota1)
ws_f.write('B17', 'El coste de lo que tienes usa el precio medio de compra. Hacienda aplica el método FIFO al vender: para el impuesto real, usa la calculadora de wealthuncut.com.', f_nota1)
ws_f.write('B18', 'Desviación en rojo: tu fondo se aleja más de 5 puntos de su peso objetivo. Revisa si toca reequilibrar.', f_nota1)
ws_f.freeze_panes(4, 2)
ws_f.set_landscape()
ws_f.fit_to_pages(1, 1)
ws_f.print_area('B1:P18')

# ------------------------------------------------------------- Valoraciones
ws_v.set_column('B:B', 14)
ws_v.set_column('C:F', 20)
ws_v.set_column('H:I', 14, None, {'hidden': True})
ws_v.write('B1', 'Evolución de tu cartera', f_titulo)
ws_v.write('B2', 'Una vez al mes (o al trimestre), apunta cuánto vale tu cartera. La plantilla calcula cuánto habías aportado y dibuja la evolución.', f_sub)
for i, h in enumerate(['Fecha', 'Valor de la cartera (€)', 'Aportado neto hasta entonces (€)', 'Ganancia (€)', 'Rentabilidad sobre lo aportado']):
    ws_v.write(3, 1 + i, h, f_h)
ws_v.set_row(3, 42)


def valor_en(a, m):
    tot = 0
    for k in range(4):
        part = 0.0
        for o in OPS:
            if o[1] == k and (o[0].year, o[0].month) <= (a, m):
                if o[2] in ('Compra', 'Traspaso entrada'):
                    part += o[3]
                else:
                    part -= o[3]
        tot += part * vl[(k, a, m)]
    return round(tot, 2)


SNAP = [(a, m) for (a, m) in meses]
for r in range(V0, V1 + 1):
    fila = r - 1
    i = r - V0
    if i < len(SNAP):
        a, m = SNAP[i]
        ws_v.write_datetime(fila, 1, datetime(a, m, 28), f_in_fecha)
        ws_v.write_number(fila, 2, valor_en(a, m), f_in_eur)
    else:
        ws_v.write_blank(fila, 1, None, f_in_fecha)
        ws_v.write_blank(fila, 2, None, f_in_eur)
    ws_v.write_formula(fila, 3, f'=IF(OR($B{r}="",$C{r}=""),"",-SUMIFS({OP["K"]},{OP["L"]},"<="&$B{r}))', f_eur)
    ws_v.write_formula(fila, 4, f'=IF(OR($B{r}="",$C{r}="",$D{r}=""),"",C{r}-D{r})', f_eur_g)
    ws_v.write_formula(fila, 5, f'=IF(OR($B{r}="",$D{r}="",N($D{r})=0),"",E{r}/D{r})', f_pct)
    ws_v.write_formula(fila, 7, f'=IF(OR($B{r}="",$C{r}=""),NA(),$B{r})', f_aux_fecha)
    ws_v.write_formula(fila, 8, f'=IF(OR($B{r}="",$C{r}=""),NA(),$C{r})', f_aux_eur)
ws_v.conditional_format(f'E{V0}:F{V1}', {'type': 'cell', 'criteria': '<', 'value': 0, 'format': F(font_color=ROJO, bold=True)})
ws_v.write('B66', 'Pon la fecha del día en que miras tu cartera (por ejemplo el último día del mes) y el valor total que te enseña tu plataforma.', f_nota1)
ws_v.freeze_panes(4, 0)
ws_v.set_landscape()
ws_v.fit_to_pages(1, 1)
ws_v.print_area(f'B1:F{V0 + 24}')

# ------------------------------------------------------------- Resumen
ws_r.set_column('B:M', 12.5)
ws_r.set_column('N:N', 3)
ws_r.set_column('O:Q', 16, f_aux)
ws_r.write('B1', 'Seguimiento de inversiones', f_titulo)
ws_r.write('B2', 'Cuánto vale tu cartera de fondos indexados, cuánto has puesto, qué rentabilidad llevas y qué te tocaría pagar a Hacienda si vendieras hoy (estimación).', f_sub)
ws_r.merge_range('B4:C4', 'Fecha de valoración', F(bold=True))
ws_r.merge_range('D4:E4', '', f_in_fecha)
ws_r.write_datetime('D4', datetime(2026, 9, 30), f_in_fecha)
ws_r.write('G4', 'Fecha del valor liquidativo que has puesto en la hoja Fondos.', f_nota1)

tarjetas = [
    ('B', 'C', 'VALOR ACTUAL', '=Fondos!I13', EUR0, VERDE, 'lo que vale hoy tu cartera'),
    ('D', 'E', 'APORTADO NETO', f'=-SUM({OP["K"]})', EUR0, VERDE, 'compras y comisiones menos ventas'),
    ('F', 'G', 'GANANCIA TOTAL', f'=B9-D9', EUR0, VERDE, 'valor actual menos aportado'),
    ('H', 'I', 'RENTABILIDAD TOTAL', '=IF(SUMIFS({0},{1},"Compra")+SUMIFS({2},{1},"Compra")=0,0,F9/(SUMIFS({0},{1},"Compra")+SUMIFS({2},{1},"Compra")))'.format(OP['H'], OP['D'], OP['G']), PCT1, VERDE, 'sobre lo que has comprado'),
    ('J', 'K', 'TIR ANUAL', '=IFERROR(XIRR(Calculo!$A$1:$A$502,Calculo!$B$1:$B$502),"n/d")', PCT1, VERDE, 'rentabilidad anualizada (XIRR)'),
    ('L', 'M', 'COSTE ANUAL TER', '=Fondos!P13', EUR0, ROJO, ''),
]
ws_r.set_row(7, 20)
ws_r.set_row(8, 34)
for a, b, nombre, formula, nf, col, nota in tarjetas:
    ws_r.merge_range(f'{a}8:{b}8', nombre, F(bold=True, font_color='#FFFFFF', bg_color=col, align='center', font_size=9))
    fm = F(bold=True, font_size=17, align='center', bg_color=CREMA, num_format=nf, border=1, border_color=BORDE)
    ws_r.merge_range(f'{a}9:{b}9', '', fm)
    ws_r.write_formula(f'{a}9', formula, fm)
    ws_r.merge_range(f'{a}10:{b}10', '', F(font_size=8, font_color=GRIS, align='center', italic=True))
    if nombre == 'COSTE ANUAL TER':
        ws_r.write_formula(f'{a}10', '=IF(B9=0,"",ROUND(L9/B9*100,2)&" % de tu cartera al año")', F(font_size=8, font_color=GRIS, align='center', italic=True))
    else:
        ws_r.write(f'{a}10', nota, F(font_size=8, font_color=GRIS, align='center', italic=True))
ws_r.conditional_format('F9', {'type': 'cell', 'criteria': '<', 'value': 0, 'format': F(font_color=ROJO)})
ws_r.conditional_format('H9', {'type': 'cell', 'criteria': '<', 'value': 0, 'format': F(font_color=ROJO)})

# Impuesto estimado
ws_r.write('B12', 'SI VENDIERAS TODO HOY (ESTIMACIÓN)', F(bold=True, font_color=VERDE, font_size=11))
filas = [
    ('Valor actual de la cartera', '=B9', EUR),
    ('Coste de lo que tienes', '=Fondos!J13', EUR),
    ('Ganancia latente', '=Fondos!K13', EUR),
    ('Impuesto estimado (IRPF del ahorro)', '=SUMPRODUCT(Listas!$I$2:$I$6,Listas!$H$2:$H$6)', EUR),
    ('Te quedarían netos', '=E13-E16', EUR),
]
for k, (n, fo, nf) in enumerate(filas):
    r = 13 + k
    ws_r.merge_range(f'B{r}:D{r}', n, f_tot if k == 4 else f_celda)
    ws_r.write_formula(f'E{r}', fo, F(border=1, border_color=BORDE, num_format=nf, bold=(k == 4), bg_color=(VERDE_CLARO if k == 4 else None) or '#FFFFFF'))
# E15 (ganancia latente) alimenta los tramos de la hoja Listas
ws_r.write('B19', 'Aproximación con el precio medio de compra,', f_nota1)
ws_r.write('B20', 'sin compensar pérdidas ni otras rentas del ahorro.', f_nota1)

# Reparto por tipo de activo (auxiliar para el donut) y cartera de un vistazo
ws_r.write('O12', 'Datos de los gráficos', F(bold=True, font_color=GRIS, font_size=9))
for k, t in enumerate(TIPOS_ACTIVO):
    ws_r.write(12 + k, 14, t, f_aux)
    ws_r.write_formula(12 + k, 15, f'=SUMIFS(Fondos!$I${F0}:$I${F1},Fondos!$D${F0}:$D${F1},O{13 + k})', f_aux_eur)

ws_r.write('H12', 'TU CARTERA DE UN VISTAZO', F(bold=True, font_color=VERDE, font_size=11))
for i, h in enumerate(['Fondo', 'Valor', 'Peso', 'Rentab.']):
    if i == 0:
        ws_r.merge_range('H13:J13', 'Fondo', f_hizq)
    else:
        ws_r.write(12, 9 + i, h, f_h)
for k in range(5):
    r = 14 + k
    fr = F0 + k
    ws_r.merge_range(f'H{r}:J{r}', '', f_celda)
    ws_r.write_formula(f'H{r}', f'=IF(Fondos!B{fr}="","",Fondos!B{fr})', f_celda)
    ws_r.write_formula(f'K{r}', f'=IF(Fondos!B{fr}="","",Fondos!I{fr})', f_eur)
    ws_r.write_formula(f'L{r}', f'=IF(Fondos!B{fr}="","",Fondos!M{fr})', f_pct)
    ws_r.write_formula(f'M{r}', f'=IF(Fondos!B{fr}="","",Fondos!L{fr})', f_pct)
ws_r.write('H19', 'Si tienes más de cinco fondos, míralos todos en la hoja Fondos.', f_nota1)

# Gráficos
c1 = wb.add_chart({'type': 'doughnut'})
c1.add_series({'name': 'Reparto por tipo de activo', 'categories': '=Resumen!$O$13:$O$16', 'values': '=Resumen!$P$13:$P$16',
               'points': [{'fill': {'color': VERDE}}, {'fill': {'color': '#4F7F6C'}}, {'fill': {'color': DORADO}}, {'fill': {'color': '#CFC8B4'}}],
               'data_labels': {'percentage': True, 'num_format': '0%;-0%;;', 'font': {'name': 'Arial', 'size': 9, 'color': '#FFFFFF', 'bold': True}}})
c1.set_title({'name': 'Reparto por tipo de activo', 'name_font': {'name': 'Arial', 'size': 11, 'bold': True, 'color': VERDE}})
c1.set_hole_size(55)
c1.set_legend({'position': 'right', 'font': {'name': 'Arial', 'size': 9}})
c1.set_size({'width': 360, 'height': 280})
c1.set_chartarea({'border': {'color': BORDE}})
ws_r.insert_chart('B21', c1, {'y_offset': 6})

c2 = wb.add_chart({'type': 'bar'})
c2.add_series({'name': 'Peso objetivo', 'categories': f'=Fondos!$B${F0}:$B${F0 + 5}', 'values': f'=Fondos!$F${F0}:$F${F0 + 5}', 'fill': {'color': '#B9C9BF'}, 'gap': 60})
c2.add_series({'name': 'Peso actual', 'categories': f'=Fondos!$B${F0}:$B${F0 + 5}', 'values': f'=Fondos!$M${F0}:$M${F0 + 5}', 'fill': {'color': VERDE}})
c2.set_title({'name': 'Peso actual frente al objetivo', 'name_font': {'name': 'Arial', 'size': 11, 'bold': True, 'color': VERDE}})
c2.set_legend({'position': 'bottom', 'font': {'name': 'Arial', 'size': 9}})
c2.set_x_axis({'num_format': '0%', 'min': 0, 'num_font': {'name': 'Arial', 'size': 8}, 'major_gridlines': {'visible': True, 'line': {'color': '#E8E3D3'}}})
c2.set_y_axis({'reverse': True, 'num_font': {'name': 'Arial', 'size': 8}})
c2.set_size({'width': 360, 'height': 280})
c2.set_chartarea({'border': {'color': BORDE}})
ws_r.insert_chart('F21', c2, {'y_offset': 6})

c3 = wb.add_chart({'type': 'scatter', 'subtype': 'straight'})
c3.add_series({'name': 'Valor de la cartera', 'categories': f"=Valoraciones!$H${V0}:$H${V1}", 'values': f"=Valoraciones!$I${V0}:$I${V1}", 'line': {'color': VERDE, 'width': 2.5}})
c3.add_series({'name': 'Aportado neto', 'categories': f"=Valoraciones!$H${V0}:$H${V1}", 'values': f"=Valoraciones!$D${V0}:$D${V1}", 'line': {'color': DORADO, 'width': 2.25, 'dash_type': 'dash'}})
c3.set_title({'name': 'Lo que has puesto frente a lo que vale', 'name_font': {'name': 'Arial', 'size': 11, 'bold': True, 'color': VERDE}})
c3.set_legend({'position': 'bottom', 'font': {'name': 'Arial', 'size': 9}})
c3.set_x_axis({'num_format': 'mmm yy', 'num_font': {'name': 'Arial', 'size': 8}, 'major_gridlines': {'visible': False}})
c3.set_y_axis({'num_format': '#,##0', 'num_font': {'name': 'Arial', 'size': 8}, 'major_gridlines': {'visible': True, 'line': {'color': '#E8E3D3'}}, 'min': 0})
c3.set_size({'width': 360, 'height': 280})
c3.set_chartarea({'border': {'color': BORDE}})
c3.show_hidden_data()
ws_r.insert_chart('J21', c3, {'y_offset': 6})
ws_r.write('B37', 'Plantilla de WealthUncut (wealthuncut.com). Información educativa, no es asesoramiento financiero ni una recomendación de inversión.', f_nota1)
ws_r.set_landscape()
ws_r.fit_to_pages(1, 1)
ws_r.print_area('B1:N37')

# ------------------------------------------------------------- Proyección
ws_p.set_column('B:B', 34)
ws_p.set_column('C:F', 18)
ws_p.set_column('H:J', 14, None, {'hidden': True})
ws_p.write('B1', 'Proyección', f_titulo)
ws_p.write('B2', '¿Cuánto podrías tener si sigues aportando? Un cálculo sencillo con una rentabilidad que eliges tú: no es una previsión.', f_sub)
ws_p.write('B4', 'Valor actual de la cartera', F(bold=True))
ws_p.write_formula('C4', '=Resumen!B9', f_eur)
ws_p.write('B5', 'Aportación mensual (€)', F(bold=True))
ws_p.write_number('C5', 200, f_in_eur)
ws_p.write('B6', 'Rentabilidad anual neta supuesta (%)', F(bold=True))
ws_p.write_number('C6', 5, f_in_pct2)
ws_p.write('D6', 'Un supuesto, ya descontado el TER. Prueba varios.', f_nota1)
ws_p.write('B7', 'Años', F(bold=True))
ws_p.write_number('C7', 20, F(border=1, border_color=BORDE, font_color=AZUL, bg_color=AMARILLO, align='center'))
ws_p.data_validation('C7', {'validate': 'integer', 'criteria': 'between', 'minimum': 1, 'maximum': 40})
ws_p.data_validation('C6', {'validate': 'decimal', 'criteria': 'between', 'minimum': -20, 'maximum': 30})
for i, h in enumerate(['Año', 'Valor estimado (€)', 'Aportado desde hoy (€)', 'Intereses (€)']):
    ws_p.write(9, 1 + i, h, f_h)
for t in range(0, 41):
    r = 11 + t
    ws_p.write_number(f'B{r}', t, f_celda_c)
    ws_p.write_formula(f'C{r}', f'=IF(B{r}>$C$7,"",$C$4*(1+$C$6/100)^B{r}+IF($C$6=0,$C$5*12*B{r},$C$5*(((1+$C$6/100)^B{r}-1)/((1+$C$6/100)^(1/12)-1))))', f_eur)
    ws_p.write_formula(f'D{r}', f'=IF(B{r}>$C$7,"",$C$4+$C$5*12*B{r})', f_eur)
    ws_p.write_formula(f'E{r}', f'=IF(B{r}>$C$7,"",C{r}-D{r})', f_eur_g)
    ws_p.write_formula(f'H{r}', f'=IF(B{r}>$C$7,NA(),B{r})', f_aux)
    ws_p.write_formula(f'I{r}', f'=IF(B{r}>$C$7,NA(),C{r})', f_aux_eur)
    ws_p.write_formula(f'J{r}', f'=IF(B{r}>$C$7,NA(),D{r})', f_aux_eur)
cp = wb.add_chart({'type': 'line'})
cp.add_series({'name': 'Valor estimado', 'categories': "='Proyección'!$H$11:$H$51", 'values': "='Proyección'!$I$11:$I$51", 'line': {'color': VERDE, 'width': 2.75}})
cp.add_series({'name': 'Aportado', 'categories': "='Proyección'!$H$11:$H$51", 'values': "='Proyección'!$J$11:$J$51", 'line': {'color': DORADO, 'width': 2.25, 'dash_type': 'dash'}})
cp.set_title({'name': 'Cuánto podría crecer tu cartera', 'name_font': {'name': 'Arial', 'size': 11, 'bold': True, 'color': VERDE}})
cp.set_legend({'position': 'bottom', 'font': {'name': 'Arial', 'size': 9}})
cp.set_x_axis({'name': 'Años', 'name_font': {'name': 'Arial', 'size': 9}, 'num_font': {'name': 'Arial', 'size': 8}})
cp.set_y_axis({'num_format': '#,##0', 'num_font': {'name': 'Arial', 'size': 8}, 'major_gridlines': {'visible': True, 'line': {'color': '#E8E3D3'}}})
cp.set_size({'width': 520, 'height': 320})
cp.set_chartarea({'border': {'color': BORDE}})
cp.show_hidden_data()
ws_p.insert_chart('G4', cp, {'x_offset': 10})
ws_p.write('B53', 'Para afinar con impuestos y comisiones usa el simulador de rentabilidad neta de wealthuncut.com.', f_nota1)
ws_p.set_landscape()
ws_p.fit_to_pages(1, 1)
ws_p.print_area('B1:N52')

# ------------------------------------------------------------- Calculo (auxiliar para la TIR)
ws_c.set_column('A:B', 16)
ws_c.write('D1', 'Hoja auxiliar para calcular la TIR (XIRR). No hace falta tocarla.', f_nota1)
MIN = f'MIN({OP["L"]})'
ws_c.write_formula('A1', f'=SUMIFS({OP["K"]},{OP["L"]},{MIN})', f_aux_eur)
ws_c.write_formula('B1', f'={MIN}', f_aux_fecha)
for k in range(R0, R1 + 1):
    r = k - R0 + 2
    ws_c.write_formula(f'A{r}', f'=IF(Operaciones!L{k}="",0,IF(Operaciones!L{k}=$B$1,0,Operaciones!K{k}))', f_aux_eur)
    ws_c.write_formula(f'B{r}', f'=IF(Operaciones!L{k}="",$B$1,Operaciones!L{k})', f_aux_fecha)
ws_c.write_formula('A502', '=Resumen!B9', f_aux_eur)
ws_c.write_formula('B502', '=Resumen!D4', f_aux_fecha)

# ------------------------------------------------------------- Instrucciones
ws_i.set_column('B:B', 4)
ws_i.set_column('C:C', 94)
ws_i.write('B1', 'Seguimiento de inversiones en fondos indexados', f_titulo)
ws_i.write('B2', 'Por WealthUncut · wealthuncut.com · Gratis para uso personal', f_sub)
ws_i.merge_range('B4:C4', 'CÓMO USARLA EN 5 PASOS', f_hizq)
pasos = [
    ('1', 'En la hoja Fondos, escribe tus fondos (nombre, ISIN, tipo de activo, TER) y el peso que te gustaría que tuviera cada uno en tu cartera. Sustituye los fondos de ejemplo.'),
    ('2', 'En Operaciones apunta cada compra, venta o traspaso: fecha, fondo, tipo, participaciones y valor liquidativo de la orden. El importe sale solo.'),
    ('3', 'Cuando quieras ver cómo vas, actualiza en Fondos el valor liquidativo de hoy de cada fondo y la fecha de valoración en Resumen.'),
    ('4', 'En Resumen verás lo que vale tu cartera, lo que has aportado, la rentabilidad, la TIR anual, el coste del TER y una estimación del impuesto si vendieras todo hoy.'),
    ('5', 'Cada mes apunta en Valoraciones el valor total para dibujar la evolución, y mira en Fondos la columna "Ajuste para reequilibrar" por si te has desviado de tus pesos.'),
]
for k, (n, t) in enumerate(pasos):
    r = 5 + k
    ws_i.write(f'B{r}', n, F(bold=True, font_color=VERDE, align='center', font_size=12))
    ws_i.write(f'C{r}', t, F(text_wrap=True))
    ws_i.set_row(r - 1, 34)
ws_i.merge_range('B11:C11', 'COLORES', f_hizq)
ws_i.write('B12', '', f_in)
ws_i.write('C12', 'Fondo amarillo con letra azul: lo rellenas tú. El resto se calcula solo.', F())
ws_i.write('B13', '', F(bg_color=VERDE_CLARO, border=1, border_color=BORDE))
ws_i.write('C13', 'Fondo verde claro: totales.', F())
ws_i.write('B14', '', F(bg_color=ROJO_CLARO, border=1, border_color=BORDE))
ws_i.write('C14', 'Rojo: pérdida, desviación grande de tu peso objetivo o aviso.', F())
ws_i.merge_range('B16:C16', 'CÓMO APUNTAR CADA OPERACIÓN', f_hizq)
ops = [
    ('Compra', 'Pon las participaciones y el valor liquidativo de la orden. Si tu plataforma cobró comisión, ponla aparte.'),
    ('Venta', 'Igual que una compra. El dinero que te llega resta de "aportado neto".'),
    ('Traspaso salida / entrada', 'Un traspaso entre fondos no tributa. Apunta una "Traspaso salida" en el fondo de origen y una "Traspaso entrada" en el de destino. En la entrada, pon en "Coste heredado" el coste de lo que traspasas (el de las participaciones que salen) para no perder la ganancia pendiente.'),
]
for k, (a, b) in enumerate(ops):
    r = 17 + k
    ws_i.write(f'B{r}', '•', F(align='center', font_color=VERDE, bold=True))
    ws_i.write(f'C{r}', f'{a}: {b}', F(text_wrap=True))
    ws_i.set_row(r - 1, 44 if k == 2 else 30)
ws_i.merge_range('B21:C21', 'LÍMITES QUE CONVIENE CONOCER', f_hizq)
lim = [
    'El impuesto de Resumen es una estimación: usa el precio medio de compra y los tramos del ahorro. Hacienda aplica FIFO al vender; para ese cálculo exacto usa la calculadora de wealthuncut.com.',
    'La rentabilidad y la TIR dependen de los valores liquidativos que escribas: si te equivocas en uno, se nota.',
    'No incluye dividendos cobrados (los fondos de acumulación no los reparten), ETF en otras divisas ni otros productos.',
    'Los fondos, las operaciones y las cifras que trae son un ejemplo ficticio. No son una recomendación de ningún producto.',
]
for k, t in enumerate(lim):
    r = 22 + k
    ws_i.write(f'B{r}', '•', F(align='center', font_color=VERDE, bold=True))
    ws_i.write(f'C{r}', t, F(text_wrap=True))
    ws_i.set_row(r - 1, 32)
ws_i.merge_range('B27:C27', 'IR A…', f_hizq)
for k, (h, txt) in enumerate((('Resumen', 'Resumen: tu panel'), ('Fondos', 'Fondos: tu cartera y el reequilibrado'), ('Operaciones', 'Operaciones: compras, ventas y traspasos'),
                              ('Valoraciones', 'Valoraciones: la evolución en el tiempo'), ('Proyección', 'Proyección: cuánto podría crecer'))):
    ws_i.write_url(f'C{28 + k}', f"internal:'{h}'!A1", f_link, txt)
ws_i.merge_range('B34:C34', 'MÁS HERRAMIENTAS GRATIS', f_hizq)
ws_i.write_url('C35', 'https://wealthuncut.com/herramientas/calculadora-impuestos-venta-fondos/', f_link, 'Impuestos al vender un fondo (con FIFO)')
ws_i.write_url('C36', 'https://wealthuncut.com/herramientas/calculadora-interes-compuesto/', f_link, 'Calculadora de interés compuesto')
ws_i.write_url('C37', 'https://wealthuncut.com/carteras/', f_link, 'Carteras: simula tu cartera de fondos')
ws_i.write_url('C38', 'https://wealthuncut.com/plantilla-seguimiento-inversiones/', f_link, 'Guía de esta plantilla, con imágenes y vídeos')
ws_i.merge_range('B40:C41', 'Información educativa, no asesoramiento financiero ni recomendación de inversión. Plantilla de David Pérez Mitjà (WealthUncut), versión 1.0 de octubre de 2026. '
                 'Puedes usarla y modificarla para ti; no la vendas ni la presentes como tuya.', f_nota)
ws_i.set_landscape()
ws_i.fit_to_pages(1, 1)
ws_i.activate()
ws_i.set_first_sheet()

wb.close()
print('Plantilla creada en', RUTA)
