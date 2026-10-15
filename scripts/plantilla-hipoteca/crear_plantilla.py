"""
Genera la plantilla de hipoteca de WealthUncut (Excel): cuadro de amortización con revisiones de tipo y
amortizaciones anticipadas.

Uso (necesita `pip install xlsxwriter openpyxl`):
    python scripts/plantilla-hipoteca/crear_plantilla.py <ruta-salida.xlsx> [xlsx-calculado-por-Excel]

El segundo argumento es opcional: un xlsx ya calculado por Excel, del que se copian los resultados de las
fórmulas como valores en caché (así las vistas previas sin Excel no muestran celdas vacías). Los datos de la
hipoteca son de ejemplo, no recomendaciones. Flujo: generar el borrador, abrirlo y guardarlo con Excel
(CalculateFull) y volver a generar pasando ese archivo como segundo argumento.
"""
import sys
from datetime import date

import xlsxwriter

RUTA = sys.argv[1] if len(sys.argv) > 1 else 'Plantilla-hipoteca-WealthUncut.xlsx'
CACHE = sys.argv[2] if len(sys.argv) > 2 else None

_valores = {}
if CACHE:
    import openpyxl
    from openpyxl.utils.datetime import to_excel
    from xlsxwriter.worksheet import Worksheet as _W

    for _ws in openpyxl.load_workbook(CACHE, data_only=True):
        for _fila in _ws.iter_rows():
            for _c in _fila:
                if _c.value is not None:
                    v = _c.value
                    if hasattr(v, 'year'):
                        v = to_excel(v)
                    _valores[(_ws.title, _c.row - 1, _c.column - 1)] = v
    _orig = _W._write_formula

    def _con_cache(self, row, col, formula, cell_format=None, value=0):
        return _orig(self, row, col, formula, cell_format, _valores.get((self.name, row, col), ''))

    _W._write_formula = _con_cache

VERDE = '#1E4A3A'
VERDE_CLARO = '#E3EDE7'
ROJO = '#B3402A'
TINTA = '#1F2A24'
GRIS = '#6B7A72'
BORDE = '#D9D2C0'
AZUL = '#0000FF'
AMARILLO = '#FFF8DC'
DORADO = '#D8A65C'

wb = xlsxwriter.Workbook(RUTA)
wb.set_properties({
    'title': 'Plantilla de hipoteca: cuadro de amortización',
    'subject': 'Cuota, intereses, revisiones de tipo y amortizaciones anticipadas',
    'author': 'David Pérez Mitjà · WealthUncut',
    'company': 'WealthUncut',
    'comments': 'Plantilla gratuita de wealthuncut.com. Información educativa, no asesoramiento financiero.',
    'website': 'https://wealthuncut.com/plantilla-hipoteca/',
})
_cache = {}


def F(**kw):
    base = {'font_name': 'Arial', 'font_size': 10, 'font_color': TINTA, 'valign': 'vcenter'}
    base.update(kw)
    k = tuple(sorted(base.items()))
    if k not in _cache:
        _cache[k] = wb.add_format(base)
    return _cache[k]


EUR = '#,##0.00 "€"'
EUR_G = '#,##0.00 "€";-#,##0.00 "€";"-"'
EUR0 = '#,##0 "€"'
PCT2 = '0.00%'
FECHA = 'dd/mm/yyyy'

f_titulo = F(font_size=20, bold=True, font_color=VERDE)
f_sub = F(font_color=GRIS, italic=True)
f_h = F(bold=True, font_color='#FFFFFF', bg_color=VERDE, align='center', border=1, border_color=VERDE, text_wrap=True)
f_hizq = F(bold=True, font_color='#FFFFFF', bg_color=VERDE, align='left', border=1, border_color=VERDE, indent=1)
f_lab = F(border=1, border_color=BORDE, indent=1)
f_celda = F(border=1, border_color=BORDE)
f_celda_c = F(border=1, border_color=BORDE, align='center')
f_eur = F(border=1, border_color=BORDE, num_format=EUR_G)
f_fecha = F(border=1, border_color=BORDE, num_format=FECHA, align='center')
f_pct = F(border=1, border_color=BORDE, num_format=PCT2, align='center')
f_in_eur = F(border=1, border_color=BORDE, font_color=AZUL, bg_color=AMARILLO, num_format=EUR)
f_in_num = F(border=1, border_color=BORDE, font_color=AZUL, bg_color=AMARILLO, num_format='0', align='center')
f_in_pct = F(border=1, border_color=BORDE, font_color=AZUL, bg_color=AMARILLO, num_format=PCT2, align='center')
f_in_fecha = F(border=1, border_color=BORDE, font_color=AZUL, bg_color=AMARILLO, num_format=FECHA, align='center')
f_in_txt = F(border=1, border_color=BORDE, font_color=AZUL, bg_color=AMARILLO, align='center')
f_kpi_l = F(border=1, border_color=BORDE, bg_color=VERDE_CLARO, indent=1)
f_kpi_eur = F(bold=True, border=1, border_color=BORDE, bg_color=VERDE_CLARO, num_format=EUR, align='right')
f_kpi_fecha = F(bold=True, border=1, border_color=BORDE, bg_color=VERDE_CLARO, num_format=FECHA, align='right')
f_kpi_txt = F(bold=True, border=1, border_color=BORDE, bg_color=VERDE_CLARO, align='right')
f_kpi_num = F(bold=True, border=1, border_color=BORDE, bg_color=VERDE_CLARO, num_format='0', align='right')
f_aux = F(font_color=GRIS, font_size=9, num_format=EUR_G)
f_nota = F(font_color=GRIS, font_size=9, italic=True, text_wrap=True, valign='top')
f_nota1 = F(font_color=GRIS, font_size=9, italic=True)


def hoja(nombre, color):
    ws = wb.add_worksheet(nombre)
    ws.set_tab_color(color)
    ws.hide_gridlines(2)
    ws.set_column('A:A', 2, F())
    ws.set_default_row(18)
    ws.set_row(0, 32)
    return ws


ws_i = hoja('Instrucciones', VERDE)
ws_h = hoja('Hipoteca', VERDE)
ws_c = hoja('Cuadro', ROJO)
ws_a = hoja('Anual', GRIS)
ws_l = hoja('Listas', BORDE)

N = 480                     # cuotas máximas (40 años)
C0, C1 = 6, 6 + N - 1       # filas del cuadro (1-indexadas)
REV0, REV1 = 4, 9           # filas de revisiones de tipo
EXT0, EXT1 = 4, 15          # filas de amortizaciones anticipadas

# ------------------------------------------------------------- Listas
ws_l.write('B1', 'Listas', f_titulo)
ws_l.write('B3', 'Modo de amortización anticipada', f_hizq)
ws_l.write('B4', 'Reducir plazo', f_celda)
ws_l.write('B5', 'Reducir cuota', f_celda)
ws_l.set_column('B:B', 34)

# ------------------------------------------------------------- Hipoteca (datos y resumen)
ws_h.set_column('B:B', 42)
ws_h.set_column('C:C', 18)
ws_h.set_column('D:D', 3)
ws_h.set_column('E:F', 16)
ws_h.set_column('G:G', 3)
ws_h.set_column('H:I', 16)
ws_h.write('B1', 'Tu hipoteca', f_titulo)
ws_h.write('B2', 'Rellena las celdas amarillas en azul. El cuadro de amortización y el resumen se calculan solos.', f_sub)

ws_h.write('B3', 'Datos del préstamo', f_hizq)
ws_h.write('C3', '', f_hizq)
datos = [
    ('Capital pendiente (€)', 180000, f_in_eur),
    ('Plazo restante (años)', 30, f_in_num),
    ('TIN anual inicial', 0.03, f_in_pct),
    ('Fecha de la primera cuota', date(2026, 11, 1), f_in_fecha),
    ('Qué haces al amortizar anticipadamente', 'Reducir plazo', f_in_txt),
    ('Comisión por amortización anticipada', 0.0, f_in_pct),
]
for i, (lab, val, fmt) in enumerate(datos):
    r = 4 + i
    ws_h.write(f'B{r}', lab, f_lab)
    if isinstance(val, date):
        ws_h.write_datetime(f'C{r}', __import__('datetime').datetime(val.year, val.month, val.day), fmt)
    else:
        ws_h.write(f'C{r}', val, fmt)
ws_h.data_validation('C4', {'validate': 'decimal', 'criteria': 'between', 'minimum': 0, 'maximum': 100000000,
                            'error_title': 'Capital', 'error_message': 'Escribe un importe en euros.'})
ws_h.data_validation('C5', {'validate': 'integer', 'criteria': 'between', 'minimum': 1, 'maximum': 40,
                            'error_title': 'Plazo', 'error_message': 'Entre 1 y 40 años.'})
ws_h.data_validation('C6', {'validate': 'decimal', 'criteria': 'between', 'minimum': 0, 'maximum': 0.25,
                            'error_title': 'TIN', 'error_message': 'Escribe el tipo como porcentaje, por ejemplo 3%.'})
ws_h.data_validation('C8', {'validate': 'list', 'source': '=Listas!$B$4:$B$5'})
ws_h.data_validation('C9', {'validate': 'decimal', 'criteria': 'between', 'minimum': 0, 'maximum': 0.05})

ws_h.write('E3', 'Revisión de tipo (fecha)', f_h)
ws_h.write('F3', 'Nuevo TIN', f_h)
revisiones = [(date(2027, 11, 1), 0.026), (date(2028, 11, 1), 0.029)]
for k, r in enumerate(range(REV0, REV1 + 1)):
    if k < len(revisiones):
        d, t = revisiones[k]
        ws_h.write_datetime(f'E{r}', __import__('datetime').datetime(d.year, d.month, d.day), f_in_fecha)
        ws_h.write(f'F{r}', t, f_in_pct)
    else:
        ws_h.write_blank(f'E{r}', None, f_in_fecha)
        ws_h.write_blank(f'F{r}', None, f_in_pct)

ws_h.write('H3', 'Amortización anticipada (fecha)', f_h)
ws_h.write('I3', 'Importe (€)', f_h)
extras = [(date(2027, 6, 1), 5000), (date(2029, 1, 1), 10000)]
for k, r in enumerate(range(EXT0, EXT1 + 1)):
    if k < len(extras):
        d, imp = extras[k]
        ws_h.write_datetime(f'H{r}', __import__('datetime').datetime(d.year, d.month, d.day), f_in_fecha)
        ws_h.write(f'I{r}', imp, f_in_eur)
    else:
        ws_h.write_blank(f'H{r}', None, f_in_fecha)
        ws_h.write_blank(f'I{r}', None, f_in_eur)
ws_h.set_row(2, 30)
ws_h.merge_range('E10:F12', 'Apunta cada cambio de tipo (por ejemplo, la revisión anual del euríbor) con la fecha desde la que rige.', f_nota)
ws_h.merge_range('H16:I17', 'Una fila por cada entrega extra. Se aplica en el mes de la fecha.', f_nota)

# Resumen
ws_h.write('B11', 'Resumen', f_hizq)
ws_h.write('C11', '', f_hizq)
rng = lambda col: f'Cuadro!${col}${C0}:${col}${C1}'
resumen = [
    ('Cuota inicial', f'=Cuadro!F{C0}', f_kpi_eur),
    ('Cuota más alta', f'=MAX({rng("F")})', f_kpi_eur),
    ('Total de intereses', f'=SUM({rng("G")})', f_kpi_eur),
    ('Total que habrás pagado', f'=SUM({rng("F")})+SUM({rng("I")})+SUM({rng("I")})*$C$9', f_kpi_eur),
    ('Cuotas que pagarás', f'=COUNTIF({rng("F")},">0")', f_kpi_num),
    ('Fecha de la última cuota', f'=IFERROR(INDEX({rng("C")},C16),"")', f_kpi_fecha),
    ('Duración', '=INT(C16/12)&" años y "&MOD(C16,12)&" meses"', f_kpi_txt),
    ('Intereses ahorrados con tus amortizaciones', f'=SUM({rng("P")})-SUM({rng("G")})', f_kpi_eur),
    ('Meses menos de hipoteca', '=$C$5*12-C16', f_kpi_num),
    ('Comisiones por amortizar', f'=SUM({rng("I")})*$C$9', f_kpi_eur),
]
for i, (lab, form, fmt) in enumerate(resumen):
    r = 12 + i
    ws_h.write(f'B{r}', lab, f_kpi_l)
    ws_h.write_formula(f'C{r}', form, fmt)

ws_h.merge_range('B23:C26', (
    'Cómo se calcula: sistema francés (cuota constante entre revisiones), con interés mensual TIN/12. En cada '
    'revisión de tipo la cuota se recalcula sobre el saldo y el plazo que quedan. "Reducir plazo" mantiene la '
    'cuota y acorta la hipoteca; "Reducir cuota" mantiene la fecha final y baja la cuota. No incluye seguros '
    'vinculados, gastos ni bonificaciones: el TAE real puede ser mayor.'), f_nota)

# ------------------------------------------------------------- Cuadro de amortización
ws_c.set_column('B:B', 6)
ws_c.set_column('C:C', 12)
ws_c.set_column('D:D', 9)
ws_c.set_column('E:J', 14)
ws_c.set_column('K:L', 9)
ws_c.set_column('M:M', 2)
ws_c.set_column('N:R', 13)
ws_c.write('B1', 'Cuadro de amortización', f_titulo)
ws_c.write('B2', 'Mes a mes. Las columnas de la derecha (en gris) son la misma hipoteca sin tus amortizaciones anticipadas, para comparar.', f_sub)
cab = ['Nº', 'Fecha', 'TIN', 'Saldo inicial', 'Cuota', 'Intereses', 'Capital', 'Amortización anticipada', 'Saldo final', 'Meses que quedan', 'Año']
for j, h in enumerate(cab):
    ws_c.write(4, 1 + j, h, f_h)
for j, h in enumerate(['Saldo inicial', 'Cuota', 'Intereses', 'Capital', 'Saldo final']):
    ws_c.write(4, 13 + j, h + ' (sin amortizar)', f_h)
ws_c.set_row(4, 40)
ws_c.freeze_panes(5, 3)

for k in range(N):
    r = C0 + k
    p = r - 1
    ws_c.write_formula(f'B{r}', f'=ROW()-{C0 - 1}', f_celda_c)
    ws_c.write_formula(f'C{r}', f'=EDATE(Hipoteca!$C$7,B{r}-1)', f_fecha)
    ws_c.write_formula(
        f'D{r}',
        f'=IFERROR(LOOKUP(2,1/((Hipoteca!$E${REV0}:$E${REV1}<>"")*(Hipoteca!$E${REV0}:$E${REV1}<=C{r})),Hipoteca!$F${REV0}:$F${REV1}),Hipoteca!$C$6)',
        f_pct)
    if k == 0:
        ws_c.write_formula(f'E{r}', '=Hipoteca!$C$4', f_eur)
        ws_c.write_formula(f'K{r}', '=Hipoteca!$C$5*12', f_celda_c)
        ws_c.write_formula(f'F{r}', f'=IF(E{r}<=0.005,0,MIN(E{r}*(1+D{r}/12),PMT(D{r}/12,K{r},-E{r})))', f_eur)
        ws_c.write_formula(f'N{r}', '=Hipoteca!$C$4', f_aux)
        ws_c.write_formula(f'O{r}', f'=IF(N{r}<=0.005,0,PMT(D{r}/12,Hipoteca!$C$5*12,-N{r}))', f_aux)
    else:
        ws_c.write_formula(f'E{r}', f'=J{p}', f_eur)
        ws_c.write_formula(
            f'K{r}',
            f'=IF(E{r}<=0.005,0,IF(Hipoteca!$C$8="Reducir cuota",$K${C0}-(B{r}-1),'
            f'IFERROR(MAX(1,ROUNDUP(NPER(D{p}/12,-F{p},E{r}),0)),$K${C0}-(B{r}-1))))',
            f_celda_c)
        ws_c.write_formula(
            f'F{r}',
            f'=IF(E{r}<=0.005,0,MIN(E{r}*(1+D{r}/12),IF(OR(Hipoteca!$C$8="Reducir cuota",D{r}<>D{p}),PMT(D{r}/12,K{r},-E{r}),F{p})))',
            f_eur)
        ws_c.write_formula(f'N{r}', f'=R{p}', f_aux)
        ws_c.write_formula(
            f'O{r}',
            f'=IF(N{r}<=0.005,0,IF(D{r}<>D{p},PMT(D{r}/12,Hipoteca!$C$5*12-(B{r}-1),-N{r}),O{p}))',
            f_aux)
    ws_c.write_formula(f'G{r}', f'=E{r}*D{r}/12', f_eur)
    ws_c.write_formula(f'H{r}', f'=IF(E{r}<=0.005,0,MIN(E{r},F{r}-G{r}))', f_eur)
    ws_c.write_formula(
        f'I{r}',
        f'=MIN(E{r}-H{r},SUMPRODUCT((YEAR(Hipoteca!$H${EXT0}:$H${EXT1})=YEAR(C{r}))*(MONTH(Hipoteca!$H${EXT0}:$H${EXT1})=MONTH(C{r}))*'
        f'(Hipoteca!$H${EXT0}:$H${EXT1}<>"")*Hipoteca!$I${EXT0}:$I${EXT1}))',
        f_eur)
    ws_c.write_formula(f'J{r}', f'=MAX(0,E{r}-H{r}-I{r})', f_eur)
    ws_c.write_formula(f'L{r}', f'=YEAR(C{r})', f_celda_c)
    ws_c.write_formula(f'P{r}', f'=N{r}*D{r}/12', f_aux)
    ws_c.write_formula(f'Q{r}', f'=IF(N{r}<=0.005,0,MIN(N{r},O{r}-P{r}))', f_aux)
    ws_c.write_formula(f'R{r}', f'=MAX(0,N{r}-Q{r})', f_aux)

# ------------------------------------------------------------- Anual
ws_a.set_column('B:B', 8)
ws_a.set_column('C:F', 16)
ws_a.write('B1', 'Año a año', f_titulo)
ws_a.write('B2', 'Cuánto pagas cada año de intereses y de capital, y cuánto debes al terminar el año.', f_sub)
for j, h in enumerate(['Año', 'Intereses', 'Capital (cuotas)', 'Amortizaciones', 'Deuda a 31 dic.']):
    ws_a.write(3, 1 + j, h, f_h)
ws_a.set_row(3, 30)
NA = 40
for k in range(NA):
    r = 5 + k
    ws_a.write_formula(f'B{r}', f'=YEAR(Hipoteca!$C$7)+{k}', f_celda_c)
    ws_a.write_formula(f'C{r}', f'=SUMIFS({rng("G")},{rng("L")},B{r})', f_eur)
    ws_a.write_formula(f'D{r}', f'=SUMIFS({rng("H")},{rng("L")},B{r})', f_eur)
    ws_a.write_formula(f'E{r}', f'=SUMIFS({rng("I")},{rng("L")},B{r})', f_eur)
    ws_a.write_formula(f'F{r}', f'=IFERROR(INDEX({rng("J")},MATCH(B{r},{rng("L")},1)),0)', f_eur)

# Gráficos
g1 = wb.add_chart({'type': 'line'})
g1.add_series({'name': 'Con tus amortizaciones', 'categories': f'=Cuadro!$C${C0}:$C${C1}', 'values': f'=Cuadro!$J${C0}:$J${C1}',
               'line': {'color': VERDE, 'width': 2.5}})
g1.add_series({'name': 'Sin amortizar', 'categories': f'=Cuadro!$C${C0}:$C${C1}', 'values': f'=Cuadro!$R${C0}:$R${C1}',
               'line': {'color': DORADO, 'width': 2, 'dash_type': 'dash'}})
g1.set_title({'name': 'Cuánto debes con el tiempo', 'name_font': {'name': 'Arial', 'size': 11, 'color': TINTA}})
g1.set_x_axis({'date_axis': True, 'num_format': 'yyyy', 'major_unit': 5, 'major_unit_type': 'years', 'num_font': {'name': 'Arial', 'size': 9}})
g1.set_y_axis({'num_format': '#,##0 "€"', 'num_font': {'name': 'Arial', 'size': 9}, 'major_gridlines': {'visible': True, 'line': {'color': BORDE}}})
g1.set_legend({'position': 'bottom', 'font': {'name': 'Arial', 'size': 9}})
g1.set_size({'width': 480, 'height': 300})
ws_h.insert_chart('B28', g1)

g2 = wb.add_chart({'type': 'column', 'subtype': 'stacked'})
g2.add_series({'name': 'Intereses', 'categories': f'=Anual!$B$5:$B${4 + NA}', 'values': f'=Anual!$C$5:$C${4 + NA}', 'fill': {'color': ROJO}})
g2.add_series({'name': 'Capital', 'categories': f'=Anual!$B$5:$B${4 + NA}', 'values': f'=Anual!$D$5:$D${4 + NA}', 'fill': {'color': VERDE}})
g2.set_title({'name': 'Qué pagas cada año: intereses y capital', 'name_font': {'name': 'Arial', 'size': 11, 'color': TINTA}})
g2.set_y_axis({'num_format': '#,##0 "€"', 'num_font': {'name': 'Arial', 'size': 9}, 'major_gridlines': {'visible': True, 'line': {'color': BORDE}}})
g2.set_x_axis({'num_font': {'name': 'Arial', 'size': 9}})
g2.set_legend({'position': 'bottom', 'font': {'name': 'Arial', 'size': 9}})
g2.set_size({'width': 440, 'height': 300})
ws_h.insert_chart('E28', g2, {'x_offset': 20})

# ------------------------------------------------------------- Instrucciones
ws_i.set_column('B:B', 4)
ws_i.set_column('C:C', 100)
ws_i.write('B1', 'Plantilla de hipoteca', f_titulo)
ws_i.write('B2', 'Calcula tu cuota, tus intereses y qué pasa si amortizas o cambia el euríbor. Cada paso, en la hoja Hipoteca.', f_sub)
pasos = [
    ('1', 'Rellena Datos del préstamo: capital pendiente, años que quedan, TIN y fecha de la próxima cuota.'),
    ('2', 'Si tu hipoteca es variable, apunta cada revisión de tipo (fecha y nuevo TIN) en la tabla de la derecha.'),
    ('3', 'Si ya has amortizado o vas a hacerlo, apunta la fecha y el importe de cada entrega extra.'),
    ('4', 'Elige si amortizar reduce la cuota o el plazo, y mira el resumen: intereses ahorrados y meses menos.'),
    ('5', 'Consulta el Cuadro de amortización (mes a mes) y la hoja Anual (cuánto pagas de intereses y de capital cada año).'),
]
for i, (n, t) in enumerate(pasos):
    ws_i.write(3 + i, 1, n, F(bold=True, font_color=VERDE, align='center'))
    ws_i.write(3 + i, 2, t, F(text_wrap=True))
ws_i.write('C10', 'Colores: celdas amarillas con texto azul = las rellenas tú; el resto se calcula solo. No borres las fórmulas.', f_nota1)
ws_i.write('C12', 'Límites: hasta 40 años y 6 revisiones de tipo y 12 amortizaciones anticipadas. Los datos de ejemplo son ficticios.', f_nota1)
ws_i.write('C13', 'No incluye seguros, gastos ni bonificaciones. Información educativa, no asesoramiento financiero.', f_nota1)
ws_i.write('C15', 'Más herramientas gratis en wealthuncut.com (simulador de hipoteca con euríbor, amortizar, alquilar o comprar).', f_nota1)

# Impresión: vertical u horizontal ajustada al ancho, para que PDF e impresión salgan limpios.
ws_i.set_portrait(); ws_i.fit_to_pages(1, 1); ws_i.print_area('A1:C16')
ws_h.set_landscape(); ws_h.fit_to_pages(1, 1); ws_h.print_area('A1:J46')
ws_c.set_landscape(); ws_c.fit_to_pages(1, 0); ws_c.print_area('A1:R45'); ws_c.repeat_rows(4)
ws_a.set_portrait(); ws_a.fit_to_pages(1, 0); ws_a.print_area('A1:F34')

wb.close()
print('Hecho:', RUTA)
